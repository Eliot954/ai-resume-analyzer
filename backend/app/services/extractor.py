from io import BytesIO
from fastapi import UploadFile


class TextExtractionError(Exception):
    pass


async def extract_text(file: UploadFile) -> str:
    content = await file.read()
    buffer = BytesIO(content)

    filename = (file.filename or "").lower()

    if filename.endswith(".pdf"):
        return _extract_pdf(buffer)
    elif filename.endswith(".docx"):
        return _extract_docx(buffer)
    elif filename.endswith(".doc"):
        return _extract_docx(buffer)
    else:
        raise TextExtractionError(f"Unsupported file extension: {filename}")


def _extract_pdf(buffer: BytesIO) -> str:
    import pdfplumber

    with pdfplumber.open(buffer) as pdf:
        texts = []
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                texts.append(page_text)

    result = "\n\n".join(texts).strip()
    if not result:
        raise TextExtractionError("Could not extract text from PDF. The file may be image-based or encrypted.")
    return result


def _extract_docx(buffer: BytesIO) -> str:
    from docx import Document

    doc = Document(buffer)
    texts = []

    for para in doc.paragraphs:
        if para.text.strip():
            texts.append(para.text.strip())

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                texts.append(row_text)

    result = "\n".join(texts).strip()
    if not result:
        raise TextExtractionError("Could not extract text from DOCX. The file may be empty or image-based.")
    return result
